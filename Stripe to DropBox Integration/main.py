import dropbox
from openpyxl import Workbook
import stripe
from io import BytesIO
from openpyxl.styles import Font
import datetime
import io
import os


def connect_spreadsheet(token):
    dbx = dropbox.Dropbox(token)
    workbook = Workbook()
    return dbx,workbook
def SettingUp(wb):
    worksheet_names = [worksheet.title for worksheet in wb.worksheets]
    sheet_names = ["Customers", "Payments", "Subscriptions", "invoice", "Events"]
    
    for sheet_name in sheet_names:
        if sheet_name not in worksheet_names:
            wb.create_sheet(title=sheet_name)
def connect_stripe(key):
    stripe.api_key = key

def customer_data():
    customers = stripe.Customer.list()
    customer_data = []
    for customer in customers.data:
        userD = {
            "Customer ID": customer.id,
            "Email Address": customer.email if customer.email else "Unknown",
            "Name": customer.name if customer.name else "Unknown",
            "Phone Number": customer.phone if customer.phone else "Unknown",
            "City": customer.address.city if customer.address and customer.address.city else "Unknown",
            "State": customer.address.state if customer.address and customer.address.state else "Unknown",
            "Country": customer.address.country if customer.address and customer.address.country else "Unknown",
            "Creation Date": format_date(customer.created) if customer.created else "Unknown"
        }
        customer_data.append(userD)
    
    return customer_data

def subscription_data():
    subscriptions = stripe.Subscription.list()
    subscription_data = []
    for subscription in subscriptions.data:
        subscription_info = {
            "Subscription ID": subscription.id,
            "Status": subscription.status,
            "Plan Name": subscription.plan.nickname if subscription.plan and subscription.plan.nickname else "Unknown",
            "Plan Amount": subscription.plan.amount if subscription.plan and subscription.plan.amount else "Unknown",
            "Billing Cycle": subscription.plan.interval if subscription.plan and subscription.plan.interval else "Unknown",
            "Start Date": format_date(subscription.start_date) if subscription.start_date else "Unknown",
            "End Date": format_date(subscription.current_period_end) if subscription.current_period_end else "Unknown",
            "Trial Period": subscription.plan.trial_period_days if subscription.plan else None,
            "Renewal Date": format_date(subscription.current_period_end) if subscription.current_period_end else "Unknown"
        }
        subscription_data.append(subscription_info)
    
    return subscription_data

def payment_data():
    payments = stripe.PaymentIntent.list()
    payment_data = []
    for payment in payments.data:
        payment_info = {
            "Payment ID": payment.id,
            "Amount": payment.amount,
            "Currency": payment.currency,
            "Payment Method": payment.payment_method if payment.payment_method else "Unknown",
            "Status": payment.status if payment.status else "Unknown",
            "Payment Date": format_date(payment.created) if payment.created else "Unknown"
        }
        payment_data.append(payment_info)
    
    return payment_data

def invoice_data():
    invoices = stripe.Invoice.list()
    invoice_data = []
    for invoice in invoices.data:
        invoice_data.append({
            "Invoice ID": invoice.id,
            "Customer ID": invoice.customer,
            "Amount Due": invoice.amount_due,
            "Amount Paid": invoice.amount_paid,
            "Due Date": invoice.due_date,
            "Status": invoice.status,
            "Invoice Date": format_date(invoice.created)
        })
    return invoice_data

def format_date(timestamp):
    date = datetime.datetime.fromtimestamp(timestamp)
    formatted_date = date.strftime('%Y-%m-%d')
    return formatted_date

def event_data():
    events = stripe.Event.list()
    event_data = []
    for event in events.data:
        event_info = {
            "Event ID": event.id,
            "Type": event.type,
            "Created Date": format_date(event.created),
            "Object Id":event.data.object.id if event.data and event.data.object else "Unknown",
            "Livemode":event.livemode,
            "Data Payload": event.data.object.object if event.data and event.data.object else "Unknown"
        }
        event_data.append(event_info)
    
    return event_data

def insert_data_from_objects(worksheet, objects):
    if len(objects) == 0:
        worksheet['A1'] = f"No {worksheet.title} available"
        return None
    
    header = list(objects[0].keys())
    num_rows = len(objects)
    num_columns = len(header)
    
    # Prepare the values for writing
    values = [header]
    
    for obj in objects:
        row_values = [obj[key] for key in header]
        values.append(row_values)
    
    # Write data to the worksheet
    for row in range(num_rows + 1):
        for col in range(num_columns):
            cell = worksheet.cell(row=row+1, column=col+1)
            cell.value = values[row][col]
    
    # Apply formatting to the header row
    header_font = Font(bold=True, size=12)
    for col in range(num_columns):
        cell = worksheet.cell(row=1, column=col+1)
        cell.font = header_font

def upload_to_dropbox(dbx, workbook, filename):
    with io.BytesIO() as f:
        workbook.save(f)
        f.seek(0)
        _, file_extension = os.path.splitext(filename)
        dropbox_filename = "/" + filename + file_extension
        dbx.files_upload(f.read(), dropbox_filename, mode=dropbox.files.WriteMode.overwrite)

def main(key,token):
    dbx,workbook=connect_spreadsheet(token)
    SettingUp(workbook)
    
    connect_stripe(key)
    customers_sheet=workbook['Customers']
    payment_sheet=workbook['Payments']
    subscriptions_sheet=workbook['Subscriptions']
    invoice_sheet=workbook['invoice']
    events_sheet=workbook['Events']
  
    insert_data_from_objects(customers_sheet,customer_data())
    insert_data_from_objects(payment_sheet,payment_data())
    insert_data_from_objects(subscriptions_sheet,subscription_data())
    insert_data_from_objects(invoice_sheet,invoice_data())
    insert_data_from_objects(events_sheet,event_data())
    upload_to_dropbox(dbx,workbook,"stripe_data.xlsx")

if __name__=="__main__":
    Stripe_api_key="sk_test_51NiymQKHVQTyd5K6ORso1XHeXEtKwG4Rr2GYHigsfQugFZiWHiRTzXHcHUdLRT1D60Q8qGilFNs0kR76qmIvplfR00AA0zVowZ"
    dropbox_token="sl.B44OQEu62bUHkMXNYZhX__ctAW5p70EYT3j_ENB6B3LbDWO9bMY9cpQDp5FIEsattUknE-N1qvukgE3zSiq53_iOwdPMxPCr9Ev89ZSIQmM936pcy0Bo5B4WFhKwnct8w1dsJfDttf9FyqBOEVkBGFw"
    main(Stripe_api_key,dropbox_token)